"""
Pesquisa Descritores - Catalogo de Teses e Dissertacoes da CAPES
Interface grafica com analise previa, filtros dinamicos e progresso detalhado.
"""

import os
import sys
import json
import math
import time
import threading
import queue
import webbrowser
import traceback
from datetime import datetime
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

import customtkinter as ctk
import tkinter as tk
from tkinter import filedialog, messagebox


# ============================================================
# CONFIGURACOES E CONSTANTES
# ============================================================

URL = "https://catalogodeteses.capes.gov.br/catalogo-teses/rest/busca"
HEADERS = {"Content-Type": "application/json", "User-Agent": "Mozilla/5.0"}

# Campos de filtro conhecidos da API CAPES (preenchidos a partir das agregacoes)
CAMPOS_PADRAO = [
    "Grau Acadêmico",
    "Ano",
    "Autor",
    "Orientador",
    "Banca",
    "Grande Àrea Conhecimento",
    "Área Conhecimento",
    "Área Avaliação",
    "Área Concentração",
    "Nome Programa",
    "Instituição",
    "Biblioteca",
]

# Configuracao persistente do app (~/.pesquisa_capes_config.json)
CONFIG_FILE = Path.home() / ".pesquisa_capes_config.json"
ZOOM_OPCOES = ["70%", "80%", "90%", "100%", "110%", "125%", "150%"]
ZOOM_PADRAO = "90%"

# customtkinter aplica auto-scaling de DPI do sistema. Para que "100%" no nosso
# dropdown represente o tamanho natural (independente de DPI), precisamos
# dividir pelo fator de DPI atual antes de chamar set_widget_scaling.
_DPI_FACTOR = None


def _detectar_dpi():
    """Detecta o fator de DPI auto-aplicado pelo customtkinter (cacheia)."""
    global _DPI_FACTOR
    if _DPI_FACTOR is not None:
        return _DPI_FACTOR
    try:
        ctk.set_widget_scaling(1.0)
        tmp = ctk.CTk()
        tmp.withdraw()
        tmp.update_idletasks()
        _DPI_FACTOR = ctk.ScalingTracker.get_widget_scaling(tmp)
        tmp.destroy()
    except Exception:
        _DPI_FACTOR = 1.0
    if not _DPI_FACTOR or _DPI_FACTOR <= 0:
        _DPI_FACTOR = 1.0
    return _DPI_FACTOR


def aplicar_zoom(zoom_str):
    """Aplica um zoom relativo independente do DPI. zoom_str ex: '90%'."""
    try:
        desired = int(zoom_str.strip("%")) / 100
    except Exception:
        return
    dpi = _detectar_dpi()
    try:
        ctk.set_widget_scaling(desired / dpi)
    except Exception:
        pass


def carregar_config():
    if CONFIG_FILE.exists():
        try:
            return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {}


def salvar_config(cfg):
    try:
        CONFIG_FILE.write_text(json.dumps(cfg, indent=2), encoding="utf-8")
    except Exception:
        pass

FILTROS_PADRAO = [
    {"campo": "Grande Àrea Conhecimento", "valor": "CIÊNCIAS HUMANAS"},
    {"campo": "Área Conhecimento", "valor": "EDUCAÇÃO"},
    {"campo": "Nome Programa", "valor": "EDUCAÇÃO"},
]

# Cores do tema (custom)
COR_PRIMARIA = "#2563EB"
COR_PRIMARIA_HOVER = "#1D4ED8"
COR_SUCESSO = "#16A34A"
COR_SUCESSO_HOVER = "#15803D"
COR_AVISO = "#D97706"
COR_PERIGO = "#DC2626"
COR_PERIGO_HOVER = "#B91C1C"
COR_NEUTRO = "#374151"
COR_FUNDO_CARD = "#1F2937"
COR_FUNDO_CARD_2 = "#111827"
COR_BORDA = "#374151"
COR_TEXTO_FRACO = "#9CA3AF"


# ============================================================
# NUCLEO - REQUESTS PARA A API DA CAPES
# ============================================================

class CapesCore:
    def __init__(self):
        self.session = requests.Session()
        self.session.headers.update(HEADERS)

    def buscar(self, termo, filtros, pagina=1, registros_por_pagina=20):
        payload = {
            "termo": termo or "",
            "filtros": filtros or [],
            "pagina": pagina,
            "registrosPorPagina": registros_por_pagina,
        }
        r = self.session.post(URL, json=payload, timeout=45)
        r.raise_for_status()
        return r.json()

    def extrair_pagina_tese(self, link):
        r = self.session.get(link, timeout=45)
        r.raise_for_status()
        soup = BeautifulSoup(r.text, "html.parser")
        palavras_el = soup.find(id="palavras")
        resumo_el = soup.find(id="resumo")
        palavras_raw = palavras_el.get_text(" ", strip=True) if palavras_el else ""
        if palavras_raw:
            palavras_raw = palavras_raw.rstrip(".").strip()
            partes = palavras_raw.split(";") if ";" in palavras_raw else palavras_raw.split(".")
            palavras = ", ".join(p.strip() for p in partes if p.strip())
        else:
            palavras = ""
        resumo = resumo_el.get_text(" ", strip=True) if resumo_el else ""
        return palavras, resumo


# ============================================================
# WORKER - THREAD DE EXECUCAO
# ============================================================

class StopRequested(Exception):
    pass


class CapesWorker(threading.Thread):
    """Roda em background e emite mensagens via queue para a GUI."""

    def __init__(self, params, q, stop_event):
        super().__init__(daemon=True)
        self.params = params
        self.q = q
        self.stop_event = stop_event
        self.core = CapesCore()

    def _check_stop(self):
        if self.stop_event.is_set():
            raise StopRequested()

    def emit(self, **kwargs):
        self.q.put(kwargs)

    def run(self):
        try:
            mode = self.params["mode"]
            if mode == "analyze":
                self._analyze()
            elif mode == "execute":
                self._execute()
        except StopRequested:
            self.emit(tipo="cancelled")
        except requests.exceptions.RequestException as e:
            self.emit(tipo="error", message=f"Erro de rede: {e}")
        except Exception as e:
            self.emit(tipo="error", message=str(e), traceback=traceback.format_exc())

    # ------- MODO ANALISAR -------
    def _analyze(self):
        termo = self.params["termo"]
        filtros = self.params["filtros"]
        self.emit(tipo="log", level="info",
                  text=f"Conectando ao Catalogo de Teses CAPES...")
        self.emit(tipo="log", level="info",
                  text=f"Termo: '{termo or '(qualquer)'}' | Filtros aplicados: {len(filtros)}")
        for f in filtros:
            self.emit(tipo="log", level="info",
                      text=f"   - {f['campo']} = {f['valor']}")
        self.emit(tipo="log", level="info", text="Enviando requisicao para a API...")
        t0 = time.time()
        data = self.core.buscar(termo, filtros, pagina=1, registros_por_pagina=1)
        dt = time.time() - t0
        self.emit(tipo="log", level="success",
                  text=f"Resposta recebida em {dt:.2f}s. Total de resultados: {data['total']}")
        self.emit(tipo="analyze_result",
                  total=data["total"],
                  agregacoes=data.get("agregacoes", []),
                  por_pagina=data.get("registrosPorPagina", 20))

    # ------- MODO EXECUTAR -------
    def _execute(self):
        termo = self.params["termo"]
        filtros = self.params["filtros"]
        keywords = self.params["keywords"]
        pasta_base = self.params["pasta_base"]
        com_resumo = self.params["com_resumo"]

        nome_pasta = termo.strip() if termo and termo.strip() else "resultados"
        nome_pasta_seguro = "".join(
            c if (c.isalnum() or c in " -_") else "_" for c in nome_pasta
        ).strip().rstrip(".")
        pasta = os.path.join(pasta_base, nome_pasta_seguro)
        os.makedirs(pasta, exist_ok=True)

        self.emit(tipo="log", level="info", text=f"Pasta de saida: {pasta}")
        self.emit(tipo="log", level="info", text="Etapa 1/3: coletando metadados das teses...")

        self._check_stop()
        data = self.core.buscar(termo, filtros, pagina=1, registros_por_pagina=20)
        total = data["total"]
        por_pagina = data["registrosPorPagina"]
        total_paginas = math.ceil(total / por_pagina) if total else 0

        self.emit(tipo="log", level="success",
                  text=f"Total da API: {total} teses, distribuidas em {total_paginas} paginas (20 por pagina)")

        resultados = []
        if total > 0:
            resultados.extend(self._processar_pagina(data, keywords))
            self.emit(tipo="progress", phase="links",
                      current=1, total=total_paginas,
                      header=f"Coletando paginas - 1/{total_paginas}",
                      detail=f"{len(resultados)} item(ns) coletado(s) ate agora")

            t0 = time.time()
            for pagina in range(2, total_paginas + 1):
                self._check_stop()
                time.sleep(0.3)
                data = self.core.buscar(termo, filtros, pagina=pagina, registros_por_pagina=20)
                resultados.extend(self._processar_pagina(data, keywords))
                eta = self._eta(t0, pagina - 1, total_paginas - 1)
                self.emit(tipo="progress", phase="links",
                          current=pagina, total=total_paginas,
                          header=f"Coletando paginas - {pagina}/{total_paginas}{eta}",
                          detail=f"{len(resultados)} item(ns) coletado(s) ate agora")

        self.emit(tipo="log", level="success",
                  text=f"Etapa 1/3 concluida: {len(resultados)} tese(s) coletada(s)" +
                       (f" (filtro de titulo aplicado: {', '.join(keywords)})" if keywords else ""))

        # Etapa 2: extrair palavras-chave e resumo
        if com_resumo and resultados:
            self.emit(tipo="log", level="info",
                      text=f"Etapa 2/3: extraindo palavras-chave e resumos de {len(resultados)} tese(s)...")
            t0 = time.time()
            sucesso = 0
            falhas = 0
            for i, item in enumerate(resultados, 1):
                self._check_stop()
                try:
                    palavras, resumo = self.core.extrair_pagina_tese(item["link"])
                    item["palavras_chave"] = palavras
                    item["resumo"] = resumo
                    sucesso += 1
                except Exception as e:
                    item["palavras_chave"] = ""
                    item["resumo"] = ""
                    falhas += 1
                    self.emit(tipo="log", level="warn",
                              text=f"Falha ao extrair tese #{i}: {str(e)[:120]}")
                eta = self._eta(t0, i, len(resultados))
                self.emit(tipo="progress", phase="extract",
                          current=i, total=len(resultados),
                          header=f"Extraindo resumos - {i}/{len(resultados)}{eta}",
                          detail=item.get("titulo") or "")
                time.sleep(0.05)
            self.emit(tipo="log", level="success",
                      text=f"Etapa 2/3 concluida: {sucesso} OK, {falhas} falha(s)")
        else:
            for item in resultados:
                item.setdefault("palavras_chave", "")
                item.setdefault("resumo", "")
            if not com_resumo:
                self.emit(tipo="log", level="info",
                          text="Etapa 2/3 ignorada (extracao de palavras-chave/resumo desativada)")

        # Etapa 3: salvar arquivos
        self.emit(tipo="log", level="info", text="Etapa 3/3: salvando arquivos...")

        nome_arq = nome_pasta_seguro or "resultados"

        json_file = os.path.join(pasta, f"{nome_arq}.json")
        with open(json_file, "w", encoding="utf-8") as f:
            json.dump(resultados, f, ensure_ascii=False, indent=2)
        self.emit(tipo="log", level="success", text=f"JSON salvo: {json_file}")

        excel_file = os.path.join(pasta, f"{nome_arq}.xlsx")
        self._gerar_excel(resultados, excel_file, com_resumo)
        self.emit(tipo="log", level="success", text=f"Excel salvo: {excel_file}")

        self.emit(tipo="log", level="success",
                  text=f"Concluido! {len(resultados)} tese(s) salvas em '{pasta}'")
        self.emit(tipo="done", pasta=pasta, total=len(resultados))

    # ------- HELPERS -------
    def _processar_pagina(self, data, keywords):
        out = []
        for item in data.get("tesesDissertacoes", []):
            titulo = item.get("titulo", "")
            if keywords:
                tl = titulo.lower()
                if not any(k in tl for k in keywords):
                    continue
            dd = item.get("dataDefesa")
            if dd:
                try:
                    dd = datetime.fromisoformat(dd.replace("Z", "")).strftime("%d/%m/%Y")
                except Exception:
                    pass
            out.append({
                "titulo": titulo,
                "autor": item.get("autor"),
                "grauAcademico": item.get("grauAcademico"),
                "instituicao": item.get("instituicao"),
                "nomePrograma": item.get("nomePrograma"),
                "municipioPrograma": item.get("municipioPrograma"),
                "biblioteca": item.get("biblioteca"),
                "dataDefesa": dd,
                "link": item.get("link"),
            })
        return out

    def _eta(self, t0, current, total):
        if not current or not total or current >= total:
            return ""
        elapsed = time.time() - t0
        if elapsed < 1:
            return ""
        rate = current / elapsed
        remaining = (total - current) / rate if rate > 0 else 0
        if remaining < 60:
            return f" (ETA {int(remaining)}s)"
        m = int(remaining // 60)
        s = int(remaining % 60)
        return f" (ETA {m}m{s:02d}s)"

    def _gerar_excel(self, data, excel_file, com_resumo):
        wb = Workbook()
        ws = wb.active
        ws.title = "dados"

        cols = [
            "titulo", "autor", "grauAcademico", "instituicao",
            "nomePrograma", "municipioPrograma", "biblioteca",
            "dataDefesa", "palavras_chave",
        ]
        if com_resumo:
            cols.append("resumo")
        cols.append("link")

        ws.append(cols)
        header_fill = PatternFill("solid", fgColor="2563EB")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", horizontal="center")

        for item in data:
            row = [item.get(c, "") for c in cols]
            ws.append(row)
            link_idx = cols.index("link") + 1
            link_cell = ws.cell(row=ws.max_row, column=link_idx)
            if link_cell.value:
                link_cell.hyperlink = link_cell.value
                link_cell.font = Font(color="0000FF", underline="single")

        widths = {
            "titulo": 60, "autor": 30, "grauAcademico": 15, "instituicao": 30,
            "nomePrograma": 30, "municipioPrograma": 20, "biblioteca": 25,
            "dataDefesa": 12, "palavras_chave": 40, "resumo": 80, "link": 40,
        }
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths.get(c, 20)
        ws.freeze_panes = "A2"

        wb.save(excel_file)


# ============================================================
# COMPONENTES DA GUI
# ============================================================

class FilterRow(ctk.CTkFrame):
    """Uma linha editavel de filtro (campo + valor)."""

    def __init__(self, parent, campos, on_remove, campo="", valor=""):
        super().__init__(parent, fg_color="transparent")
        self.on_remove = on_remove

        self.campo_var = tk.StringVar(value=campo)
        self.valor_var = tk.StringVar(value=valor)

        self.campo_combo = ctk.CTkComboBox(
            self, values=campos, variable=self.campo_var, width=230,
            button_color=COR_PRIMARIA, button_hover_color=COR_PRIMARIA_HOVER,
            dropdown_fg_color=COR_FUNDO_CARD_2,
        )
        self.campo_combo.grid(row=0, column=0, padx=(0, 8), sticky="w")

        self.valor_entry = ctk.CTkEntry(
            self, textvariable=self.valor_var,
            placeholder_text="Valor (ex: CIÊNCIAS HUMANAS)", width=320,
        )
        self.valor_entry.grid(row=0, column=1, padx=(0, 8), sticky="ew")

        self.remove_btn = ctk.CTkButton(
            self, text="✕", width=32, height=28,
            fg_color=COR_PERIGO, hover_color=COR_PERIGO_HOVER,
            command=self._remove,
        )
        self.remove_btn.grid(row=0, column=2)

        self.grid_columnconfigure(1, weight=1)

    def _remove(self):
        self.on_remove(self)

    def get(self):
        c = self.campo_var.get().strip()
        v = self.valor_var.get().strip()
        if not c or not v:
            return None
        return {"campo": c, "valor": v}

    def set_campos(self, campos):
        self.campo_combo.configure(values=campos)


# ============================================================
# APLICACAO PRINCIPAL
# ============================================================

ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")


class App(ctk.CTk):
    def __init__(self):
        # Carrega config salvo e aplica zoom antes de criar widgets
        self._config = carregar_config()
        zoom_str = self._config.get("zoom", ZOOM_PADRAO)
        if zoom_str not in ZOOM_OPCOES:
            zoom_str = ZOOM_PADRAO
        aplicar_zoom(zoom_str)

        super().__init__()
        self.title("Pesquisa Descritores - Catalogo CAPES")
        self.geometry("1280x860")
        self.minsize(900, 700)
        self._zoom_atual = zoom_str

        self.queue = queue.Queue()
        self.stop_event = threading.Event()
        self.worker = None
        self.campos_disponiveis = list(CAMPOS_PADRAO)
        self.filter_rows = []
        self.last_total = None
        self.last_pasta = None
        self.last_agregacoes = []

        # Lista de labels com wraplength dinamico: (label, container, padding)
        self._wrap_bindings = []

        self._build_ui()

        # Atualiza wraplengths sempre que a janela redimensionar
        # add="+" para nao sobrescrever bindings internos do customtkinter
        self.bind("<Configure>", self._on_root_configure, add="+")
        # Forca atualizacao inicial depois do primeiro layout
        self.after(50, self._update_all_wraplengths)
        self.after(300, self._update_all_wraplengths)

        # Maximiza a janela apos o widget estar mapeado.
        # Chamar state("zoomed") direto no __init__ as vezes nao funciona
        # porque a janela ainda nao foi exibida.
        self.after(10, self._maximize)
        # Reforco caso o sistema tenha restaurado durante o build
        self.after(200, self._maximize)

        self._tick()

    def _maximize(self):
        try:
            self.state("zoomed")
        except Exception:
            pass

    def _on_zoom_changed(self, value):
        if value not in ZOOM_OPCOES:
            return
        self._zoom_atual = value
        self._config["zoom"] = value
        salvar_config(self._config)
        # Esconde a janela durante o reescalonamento para evitar artefatos visuais
        self.withdraw()
        aplicar_zoom(value)
        self.update_idletasks()
        self.after(120, self._restaurar_apos_zoom)

    def _restaurar_apos_zoom(self):
        self.deiconify()
        self.state("zoomed")
        self.after(60, self._update_all_wraplengths)
        self.after(250, self._update_all_wraplengths)

    # ------------------------------------------------------------
    # CONSTRUCAO DA UI
    # ------------------------------------------------------------
    def _build_ui(self):
        # Layout: 2 colunas
        self.grid_columnconfigure(0, weight=1, minsize=520)
        self.grid_columnconfigure(1, weight=1, minsize=560)
        self.grid_rowconfigure(1, weight=1)

        # ------- HEADER -------
        header = ctk.CTkFrame(self, fg_color=COR_FUNDO_CARD_2, corner_radius=0, height=70)
        header.grid(row=0, column=0, columnspan=2, sticky="ew")
        header.grid_propagate(False)
        header.grid_columnconfigure(0, weight=1)

        title = ctk.CTkLabel(
            header, text="Pesquisa de Descritores",
            font=ctk.CTkFont(size=22, weight="bold"),
        )
        title.grid(row=0, column=0, sticky="w", padx=20, pady=(12, 0))

        sub = ctk.CTkLabel(
            header, text="Catalogo de Teses e Dissertacoes da CAPES",
            font=ctk.CTkFont(size=12),
            text_color=COR_TEXTO_FRACO,
        )
        sub.grid(row=1, column=0, sticky="w", padx=20, pady=(0, 12))

        # Controles do header (zoom + link CAPES)
        controles = ctk.CTkFrame(header, fg_color="transparent")
        controles.grid(row=0, column=1, rowspan=2, padx=20, pady=12, sticky="e")

        zoom_lbl = ctk.CTkLabel(
            controles, text="Zoom:", font=ctk.CTkFont(size=11),
            text_color=COR_TEXTO_FRACO,
        )
        zoom_lbl.pack(side="left", padx=(0, 6))

        self.zoom_var = tk.StringVar(value=self._zoom_atual)
        zoom_combo = ctk.CTkComboBox(
            controles, values=ZOOM_OPCOES, variable=self.zoom_var,
            width=80, height=28,
            command=self._on_zoom_changed,
            button_color=COR_PRIMARIA, button_hover_color=COR_PRIMARIA_HOVER,
            dropdown_fg_color=COR_FUNDO_CARD_2,
            state="readonly",
        )
        zoom_combo.pack(side="left", padx=(0, 12))

        link_btn = ctk.CTkButton(
            controles, text="Abrir site CAPES", width=140,
            fg_color="transparent", border_width=1, border_color=COR_PRIMARIA,
            hover_color=COR_NEUTRO,
            command=lambda: webbrowser.open("https://catalogodeteses.capes.gov.br"),
        )
        link_btn.pack(side="left")

        # ------- COLUNA ESQUERDA: CONFIGURACAO -------
        left = ctk.CTkScrollableFrame(self, fg_color="transparent", label_text="")
        left.grid(row=1, column=0, sticky="nsew", padx=(16, 8), pady=12)
        left.grid_columnconfigure(0, weight=1)
        self.left_scroll = left  # exposto para automacao

        self._build_termo(left, 0)
        self._build_filtros(left, 1)
        self._build_keywords(left, 2)
        self._build_output(left, 3)
        self._build_opcoes(left, 4)
        self._build_botoes(left, 5)

        # ------- COLUNA DIREITA: RESULTADO/PROGRESSO/LOG -------
        right = ctk.CTkFrame(self, fg_color="transparent")
        right.grid(row=1, column=1, sticky="nsew", padx=(8, 16), pady=12)
        right.grid_columnconfigure(0, weight=1)
        right.grid_rowconfigure(2, weight=1)

        self._build_analise(right, 0)
        self._build_progresso(right, 1)
        self._build_log(right, 2)

        # ------- STATUS BAR -------
        statusbar = ctk.CTkFrame(self, fg_color=COR_FUNDO_CARD_2, height=30, corner_radius=0)
        statusbar.grid(row=2, column=0, columnspan=2, sticky="ew")
        statusbar.grid_propagate(False)
        statusbar.grid_columnconfigure(0, weight=1)

        self.status_label = ctk.CTkLabel(
            statusbar, text="Pronto. Configure os filtros e clique em ANALISAR.",
            font=ctk.CTkFont(size=11), text_color=COR_TEXTO_FRACO, anchor="w",
        )
        self.status_label.grid(row=0, column=0, sticky="ew", padx=12, pady=4)

    # ----- secoes -----
    def _section(self, parent, row, title, descricao=None):
        card = ctk.CTkFrame(parent, fg_color=COR_FUNDO_CARD, corner_radius=10)
        card.grid(row=row, column=0, sticky="ew", pady=(0, 12))
        card.grid_columnconfigure(0, weight=1)
        lbl = ctk.CTkLabel(
            card, text=title, font=ctk.CTkFont(size=14, weight="bold"), anchor="w",
        )
        lbl.grid(row=0, column=0, sticky="ew", padx=14, pady=(12, 2))
        if descricao:
            sub = ctk.CTkLabel(
                card, text=descricao, font=ctk.CTkFont(size=11),
                text_color=COR_TEXTO_FRACO, anchor="w", justify="left",
                wraplength=480,
            )
            sub.grid(row=1, column=0, sticky="ew", padx=14, pady=(0, 6))
            self._bind_dynamic_wrap(card, sub, padding=36)
        body = ctk.CTkFrame(card, fg_color="transparent")
        body.grid(row=2, column=0, sticky="ew", padx=14, pady=(4, 14))
        body.grid_columnconfigure(0, weight=1)
        return body

    def _bind_dynamic_wrap(self, container, label, padding=36):
        """Registra um label cujo wraplength acompanha a largura do container."""
        self._wrap_bindings.append((label, container, padding))

    def _on_root_configure(self, event):
        # filtra: so processa quando o evento e da propria janela
        if event.widget is self:
            self._update_all_wraplengths()

    def _update_all_wraplengths(self):
        # customtkinter multiplica wraplength pelo widget_scaling antes de passar
        # ao tk widget (HiDPI). Para que o texto caiba no container, precisamos
        # dividir o wraplength logico desejado pelo scaling atual.
        try:
            scaling = ctk.ScalingTracker.get_widget_scaling(self)
        except Exception:
            scaling = 1.0
        if not scaling or scaling <= 0:
            scaling = 1.0
        for label, container, padding in self._wrap_bindings:
            try:
                w = container.winfo_width()
            except Exception:
                continue
            if w <= 1:
                continue
            new_wrap = max(180, int((w - padding) / scaling))
            try:
                current = int(label.cget("wraplength") or 0)
            except Exception:
                current = 0
            if abs(new_wrap - current) > 4:
                try:
                    label.configure(wraplength=new_wrap)
                except Exception:
                    pass

    def _build_termo(self, parent, row):
        body = self._section(
            parent, row, "1. Termo de busca",
            "Palavra ou expressao para pesquisar no catalogo da CAPES. "
            "Deixe em branco para listar todas as teses (combinado com os filtros abaixo).",
        )
        self.termo_var = tk.StringVar(value="gerencialismo")
        self.termo_entry = ctk.CTkEntry(
            body, textvariable=self.termo_var, height=36,
            placeholder_text="ex: gerencialismo, performatividade docente",
            font=ctk.CTkFont(size=13),
        )
        self.termo_entry.grid(row=0, column=0, sticky="ew")

    def _build_filtros(self, parent, row):
        body = self._section(
            parent, row, "2. Filtros",
            "Cada filtro restringe a busca por um campo (Area, Programa, Ano etc). "
            "Para o mesmo campo, o catalogo trata como OU (qualquer um dos valores). "
            "Sugestoes aparecem na coluna direita apos clicar ANALISAR.",
        )
        body.grid_columnconfigure(0, weight=1)

        self.filtros_container = ctk.CTkFrame(body, fg_color="transparent")
        self.filtros_container.grid(row=0, column=0, sticky="ew")
        self.filtros_container.grid_columnconfigure(0, weight=1)

        botoes = ctk.CTkFrame(body, fg_color="transparent")
        botoes.grid(row=1, column=0, sticky="ew", pady=(8, 0))

        add_btn = ctk.CTkButton(
            botoes, text="+ Adicionar filtro", width=160,
            fg_color=COR_NEUTRO, hover_color="#4B5563",
            command=lambda: self._add_filter_row(),
        )
        add_btn.pack(side="left")

        clear_btn = ctk.CTkButton(
            botoes, text="Limpar filtros", width=140,
            fg_color="transparent", border_width=1, border_color=COR_BORDA,
            hover_color=COR_NEUTRO, command=self._clear_filters,
        )
        clear_btn.pack(side="left", padx=8)

        preset_btn = ctk.CTkButton(
            botoes, text="Preset Educacao", width=140,
            fg_color="transparent", border_width=1, border_color=COR_BORDA,
            hover_color=COR_NEUTRO, command=self._preset_educacao,
        )
        preset_btn.pack(side="left")

        # Linhas iniciais com filtros padrao
        for f in FILTROS_PADRAO:
            self._add_filter_row(f["campo"], f["valor"])

    def _build_keywords(self, parent, row):
        body = self._section(
            parent, row, "3. Filtrar por palavras no titulo (opcional)",
            "Apos a busca, mantem apenas teses cujo titulo contenha pelo menos uma das palavras abaixo. "
            "Separe por virgula. Deixe em branco para nao filtrar.",
        )
        self.keywords_var = tk.StringVar(value="")
        self.keywords_entry = ctk.CTkEntry(
            body, textvariable=self.keywords_var, height=36,
            placeholder_text="ex: escola, docente, basica",
        )
        self.keywords_entry.grid(row=0, column=0, sticky="ew")

    def _build_output(self, parent, row):
        body = self._section(
            parent, row, "4. Pasta de saida",
            "Sera criada uma subpasta com o nome do termo dentro da pasta escolhida.",
        )
        body.grid_columnconfigure(0, weight=1)

        self.pasta_var = tk.StringVar(value=os.path.abspath(os.getcwd()))
        self.pasta_entry = ctk.CTkEntry(body, textvariable=self.pasta_var, height=36)
        self.pasta_entry.grid(row=0, column=0, sticky="ew", padx=(0, 8))

        choose_btn = ctk.CTkButton(
            body, text="Escolher...", width=110,
            fg_color=COR_NEUTRO, hover_color="#4B5563",
            command=self._choose_folder,
        )
        choose_btn.grid(row=0, column=1)

    def _build_opcoes(self, parent, row):
        body = self._section(
            parent, row, "5. Opcoes",
            "A extracao de resumo abre a pagina de cada tese individualmente "
            "e e mais lenta (~1s por tese). Desligue se quiser apenas a lista basica.",
        )
        self.com_resumo_var = tk.BooleanVar(value=True)
        cb = ctk.CTkCheckBox(
            body, text="Extrair palavras-chave e resumo de cada tese",
            variable=self.com_resumo_var,
            font=ctk.CTkFont(size=12),
        )
        cb.grid(row=0, column=0, sticky="w")

    def _build_botoes(self, parent, row):
        card = ctk.CTkFrame(parent, fg_color=COR_FUNDO_CARD_2, corner_radius=10)
        card.grid(row=row, column=0, sticky="ew", pady=(0, 16))
        card.grid_columnconfigure((0, 1, 2), weight=1)

        self.btn_analisar = ctk.CTkButton(
            card, text="🔍  ANALISAR", height=48,
            fg_color=COR_PRIMARIA, hover_color=COR_PRIMARIA_HOVER,
            font=ctk.CTkFont(size=14, weight="bold"),
            command=self._on_analisar,
        )
        self.btn_analisar.grid(row=0, column=0, padx=10, pady=14, sticky="ew")

        self.btn_executar = ctk.CTkButton(
            card, text="▶  EXECUTAR", height=48,
            fg_color=COR_SUCESSO, hover_color=COR_SUCESSO_HOVER,
            font=ctk.CTkFont(size=14, weight="bold"),
            state="disabled",
            command=self._on_executar,
        )
        self.btn_executar.grid(row=0, column=1, padx=10, pady=14, sticky="ew")

        self.btn_parar = ctk.CTkButton(
            card, text="■  PARAR", height=48,
            fg_color=COR_PERIGO, hover_color=COR_PERIGO_HOVER,
            font=ctk.CTkFont(size=14, weight="bold"),
            state="disabled",
            command=self._on_parar,
        )
        self.btn_parar.grid(row=0, column=2, padx=10, pady=14, sticky="ew")

    # ----- coluna direita -----
    def _build_analise(self, parent, row):
        card = ctk.CTkFrame(parent, fg_color=COR_FUNDO_CARD, corner_radius=10)
        card.grid(row=row, column=0, sticky="ew", pady=(0, 12))
        card.grid_columnconfigure(0, weight=1)

        title = ctk.CTkLabel(
            card, text="Resultado da analise",
            font=ctk.CTkFont(size=14, weight="bold"), anchor="w",
        )
        title.grid(row=0, column=0, sticky="ew", padx=14, pady=(12, 0))

        self.total_label = ctk.CTkLabel(
            card,
            text="Clique em ANALISAR para descobrir quantas teses\nseu filtro retorna antes de executar a coleta.",
            font=ctk.CTkFont(size=12),
            text_color=COR_TEXTO_FRACO, justify="left", anchor="w",
            wraplength=480,
        )
        self.total_label.grid(row=1, column=0, sticky="ew", padx=14, pady=(4, 8))
        self._bind_dynamic_wrap(card, self.total_label, padding=36)

        # Sugestoes
        self.agregacoes_label = ctk.CTkLabel(
            card, text="Sugestoes de filtros (clique em + para adicionar):",
            font=ctk.CTkFont(size=11, weight="bold"),
            text_color=COR_TEXTO_FRACO, anchor="w", justify="left",
            wraplength=480,
        )
        self.agregacoes_label.grid(row=2, column=0, sticky="ew", padx=14, pady=(2, 4))
        self.agregacoes_label.grid_remove()
        self._bind_dynamic_wrap(card, self.agregacoes_label, padding=36)

        self.agregacoes_frame = ctk.CTkScrollableFrame(
            card, fg_color=COR_FUNDO_CARD_2, height=180, corner_radius=8,
        )
        self.agregacoes_frame.grid(row=3, column=0, sticky="ew", padx=14, pady=(0, 14))
        self.agregacoes_frame.grid_columnconfigure(0, weight=1)
        self.agregacoes_frame.grid_remove()

    def _build_progresso(self, parent, row):
        card = ctk.CTkFrame(parent, fg_color=COR_FUNDO_CARD, corner_radius=10, height=140)
        card.grid(row=row, column=0, sticky="ew", pady=(0, 12))
        card.grid_columnconfigure(0, weight=1)
        card.grid_propagate(False)  # nao deixa o conteudo redimensionar o card

        title = ctk.CTkLabel(
            card, text="Progresso",
            font=ctk.CTkFont(size=14, weight="bold"), anchor="w",
        )
        title.grid(row=0, column=0, sticky="ew", padx=14, pady=(12, 6))

        self.progress = ctk.CTkProgressBar(card, height=18, progress_color=COR_PRIMARIA)
        self.progress.grid(row=1, column=0, sticky="ew", padx=14, pady=(0, 8))
        self.progress.set(0)

        # Linha 1 do status: sempre curta (% e contagem). Nunca quebra.
        self.progress_status = ctk.CTkLabel(
            card, text="Aguardando inicio...",
            font=ctk.CTkFont(size=12, weight="bold"),
            text_color="#E5E7EB", anchor="w", justify="left",
        )
        self.progress_status.grid(row=2, column=0, sticky="ew", padx=14, pady=(0, 2))

        # Linha 2 do status: detalhe truncado (titulo da tese, etc).
        self.progress_detail = ctk.CTkLabel(
            card, text="",
            font=ctk.CTkFont(size=11),
            text_color=COR_TEXTO_FRACO, anchor="w", justify="left",
        )
        self.progress_detail.grid(row=3, column=0, sticky="ew", padx=14, pady=(0, 12))

    def _build_log(self, parent, row):
        card = ctk.CTkFrame(parent, fg_color=COR_FUNDO_CARD, corner_radius=10)
        card.grid(row=row, column=0, sticky="nsew")
        card.grid_columnconfigure(0, weight=1)
        card.grid_rowconfigure(1, weight=1)

        head = ctk.CTkFrame(card, fg_color="transparent")
        head.grid(row=0, column=0, sticky="ew", padx=14, pady=(12, 4))
        head.grid_columnconfigure(0, weight=1)

        title = ctk.CTkLabel(
            head, text="Log detalhado",
            font=ctk.CTkFont(size=14, weight="bold"), anchor="w",
        )
        title.grid(row=0, column=0, sticky="w")

        clear_btn = ctk.CTkButton(
            head, text="Limpar", width=80, height=26,
            fg_color="transparent", border_width=1, border_color=COR_BORDA,
            hover_color=COR_NEUTRO, command=self._clear_log,
        )
        clear_btn.grid(row=0, column=1, sticky="e")

        # Usar tk.Text para suporte a tags coloridas
        log_frame = ctk.CTkFrame(card, fg_color=COR_FUNDO_CARD_2, corner_radius=8)
        log_frame.grid(row=1, column=0, sticky="nsew", padx=14, pady=(4, 14))
        log_frame.grid_rowconfigure(0, weight=1)
        log_frame.grid_columnconfigure(0, weight=1)

        self.log_text = tk.Text(
            log_frame, wrap="word",
            bg="#0B1220", fg="#E5E7EB",
            insertbackground="#E5E7EB",
            relief="flat", borderwidth=0,
            font=("Consolas", 10), padx=10, pady=10,
            state="disabled",
        )
        self.log_text.grid(row=0, column=0, sticky="nsew")

        scroll = ctk.CTkScrollbar(log_frame, command=self.log_text.yview)
        scroll.grid(row=0, column=1, sticky="ns")
        self.log_text.configure(yscrollcommand=scroll.set)

        self.log_text.tag_configure("info", foreground="#93C5FD")
        self.log_text.tag_configure("success", foreground="#86EFAC")
        self.log_text.tag_configure("warn", foreground="#FCD34D")
        self.log_text.tag_configure("error", foreground="#FCA5A5")
        self.log_text.tag_configure("ts", foreground="#6B7280")

    # ------------------------------------------------------------
    # ACOES E HANDLERS
    # ------------------------------------------------------------
    def _add_filter_row(self, campo="", valor=""):
        row = FilterRow(self.filtros_container, self.campos_disponiveis,
                        self._remove_filter_row, campo=campo, valor=valor)
        row.grid(row=len(self.filter_rows), column=0, sticky="ew", pady=4)
        self.filter_rows.append(row)

    def _remove_filter_row(self, row):
        if row in self.filter_rows:
            self.filter_rows.remove(row)
            row.destroy()
            for i, r in enumerate(self.filter_rows):
                r.grid_configure(row=i)

    def _clear_filters(self):
        for r in list(self.filter_rows):
            r.destroy()
        self.filter_rows.clear()

    def _preset_educacao(self):
        self._clear_filters()
        for f in FILTROS_PADRAO:
            self._add_filter_row(f["campo"], f["valor"])

    def _choose_folder(self):
        atual = self.pasta_var.get()
        if not os.path.isdir(atual):
            atual = os.path.expanduser("~")
        d = filedialog.askdirectory(initialdir=atual, title="Escolha a pasta de saida")
        if d:
            self.pasta_var.set(d)

    def _coletar_filtros(self):
        out = []
        for r in self.filter_rows:
            f = r.get()
            if f:
                out.append(f)
        return out

    def _coletar_keywords(self):
        raw = self.keywords_var.get().strip()
        if not raw:
            return []
        return [k.strip().lower() for k in raw.split(",") if k.strip()]

    # ------- Botoes -------
    def _on_analisar(self):
        if self.worker and self.worker.is_alive():
            return
        termo = self.termo_var.get().strip()
        filtros = self._coletar_filtros()
        params = {"mode": "analyze", "termo": termo, "filtros": filtros}
        self._start_worker(params, status="Analisando consulta...")
        self.btn_executar.configure(state="disabled")

    def _on_executar(self):
        if self.worker and self.worker.is_alive():
            return
        if self.last_total is None:
            messagebox.showwarning(
                "Analise obrigatoria",
                "Clique em ANALISAR antes de executar para confirmar quantos resultados serao processados.",
            )
            return
        if self.last_total == 0:
            messagebox.showinfo(
                "Sem resultados",
                "A consulta retorna 0 resultados. Ajuste o termo ou os filtros.",
            )
            return
        if self.last_total > 1000:
            ok = messagebox.askyesno(
                "Confirmar coleta grande",
                f"A consulta retornara {self.last_total} teses. "
                f"Pode levar varios minutos. Continuar?",
            )
            if not ok:
                return

        pasta_base = self.pasta_var.get().strip()
        if not pasta_base:
            messagebox.showerror("Pasta invalida", "Escolha uma pasta de saida.")
            return
        if not os.path.isdir(pasta_base):
            try:
                os.makedirs(pasta_base, exist_ok=True)
            except Exception as e:
                messagebox.showerror("Pasta invalida", f"Nao foi possivel criar: {e}")
                return

        params = {
            "mode": "execute",
            "termo": self.termo_var.get().strip(),
            "filtros": self._coletar_filtros(),
            "keywords": self._coletar_keywords(),
            "pasta_base": pasta_base,
            "com_resumo": self.com_resumo_var.get(),
        }
        self._start_worker(params, status="Executando coleta...")

    def _on_parar(self):
        if self.worker and self.worker.is_alive():
            self.stop_event.set()
            self._set_status("Cancelando... aguardando finalizacao.")
            self._log("Cancelamento solicitado.", "warn")

    def _start_worker(self, params, status=""):
        self.stop_event.clear()
        self.progress.set(0)
        self.progress_status.configure(text="Iniciando...")
        self.progress_detail.configure(text="")
        self._set_status(status)
        self._set_busy(True)
        self.worker = CapesWorker(params, self.queue, self.stop_event)
        self.worker.start()

    def _set_busy(self, busy):
        if busy:
            self.btn_analisar.configure(state="disabled")
            self.btn_executar.configure(state="disabled")
            self.btn_parar.configure(state="normal")
        else:
            self.btn_analisar.configure(state="normal")
            self.btn_parar.configure(state="disabled")
            if self.last_total is not None and self.last_total > 0:
                self.btn_executar.configure(state="normal")

    # ------------------------------------------------------------
    # CONSUMO DA QUEUE (loop da UI)
    # ------------------------------------------------------------
    def _tick(self):
        try:
            while True:
                msg = self.queue.get_nowait()
                self._handle_msg(msg)
        except queue.Empty:
            pass
        self.after(80, self._tick)

    def _handle_msg(self, msg):
        tipo = msg.get("tipo")
        if tipo == "log":
            self._log(msg["text"], msg.get("level", "info"))
        elif tipo == "progress":
            current = msg["current"]
            total = msg["total"]
            frac = current / total if total else 0
            self.progress.set(frac)
            header = msg.get("header") or msg.get("msg", "")
            detail = msg.get("detail", "")
            self.progress_status.configure(text=f"{int(frac * 100)}%  •  {header}")
            self.progress_detail.configure(text=self._truncate(detail, 90))
        elif tipo == "analyze_result":
            self._show_analyze_result(msg["total"], msg["agregacoes"])
            self._set_busy(False)
        elif tipo == "done":
            self.progress.set(1.0)
            self._show_done_dialog(msg["pasta"], msg["total"])
            self._set_busy(False)
        elif tipo == "cancelled":
            self._log("Operacao cancelada pelo usuario.", "warn")
            self.progress_status.configure(text="Cancelado.")
            self.progress_detail.configure(text="")
            self._set_busy(False)
            self._set_status("Cancelado.")
        elif tipo == "error":
            self._log("ERRO: " + msg["message"], "error")
            if msg.get("traceback"):
                self._log(msg["traceback"], "error")
            self.progress_status.configure(text="Erro.")
            self.progress_detail.configure(text="")
            self._set_busy(False)
            self._set_status("Erro durante a execucao. Veja o log.")
            messagebox.showerror("Erro", msg["message"])

    # ------- helpers UI -------
    def _show_analyze_result(self, total, agregacoes):
        self.last_total = total
        self.last_agregacoes = agregacoes

        cor = COR_SUCESSO if total > 0 else COR_AVISO
        emoji = "✓" if total > 0 else "⚠"
        paginas = math.ceil(total / 20) if total else 0

        self.total_label.configure(
            text=f"{emoji}  {total:,} teses encontradas\n"
                 f"   ({paginas} paginas de 20 a serem percorridas)".replace(",", "."),
            text_color=cor,
            font=ctk.CTkFont(size=14, weight="bold"),
            justify="left",
        )

        if total > 0:
            self.btn_executar.configure(state="normal")
        self._set_status(f"Analise concluida. {total} resultado(s).")

        # Atualizar campos disponiveis com base nas agregacoes retornadas
        nomes_campos = [a["campo"] for a in agregacoes]
        if nomes_campos:
            self.campos_disponiveis = nomes_campos
            for r in self.filter_rows:
                r.set_campos(self.campos_disponiveis)

        # Mostrar agregacoes
        for w in self.agregacoes_frame.winfo_children():
            w.destroy()

        if agregacoes:
            self.agregacoes_label.grid()
            self.agregacoes_frame.grid()
            self._render_agregacoes(agregacoes)
        else:
            self.agregacoes_label.grid_remove()
            self.agregacoes_frame.grid_remove()

    def _render_agregacoes(self, agregacoes):
        ordem_pref = ["Grande Àrea Conhecimento", "Área Conhecimento",
                      "Nome Programa", "Ano", "Grau Acadêmico",
                      "Instituição"]
        agregacoes_ord = sorted(
            agregacoes,
            key=lambda a: ordem_pref.index(a["campo"]) if a["campo"] in ordem_pref else 99,
        )

        row_idx = 0
        for ag in agregacoes_ord:
            campo = ag["campo"]
            agregados = ag.get("agregados", [])
            if not agregados:
                continue

            head = ctk.CTkLabel(
                self.agregacoes_frame, text=f"{campo}",
                font=ctk.CTkFont(size=12, weight="bold"),
                anchor="w", justify="left", wraplength=440,
            )
            head.grid(row=row_idx, column=0, sticky="ew", padx=8, pady=(8, 2))
            row_idx += 1

            for v in agregados[:6]:
                line = ctk.CTkFrame(self.agregacoes_frame, fg_color="transparent")
                line.grid(row=row_idx, column=0, sticky="ew", padx=8, pady=1)
                line.grid_columnconfigure(0, weight=1)

                txt = f"{v['valor']}"
                cnt = f"{v['total']:,}".replace(",", ".")
                lbl = ctk.CTkLabel(
                    line, text=f"  {txt}  -  {cnt}",
                    font=ctk.CTkFont(size=11),
                    text_color="#D1D5DB", anchor="w", justify="left",
                    wraplength=460,
                )
                lbl.grid(row=0, column=0, sticky="ew")

                btn = ctk.CTkButton(
                    line, text="+", width=28, height=22,
                    fg_color=COR_PRIMARIA, hover_color=COR_PRIMARIA_HOVER,
                    command=lambda c=campo, vv=v["valor"]: self._adicionar_filtro_da_agregacao(c, vv),
                )
                btn.grid(row=0, column=1, padx=(4, 0))
                row_idx += 1

    def _adicionar_filtro_da_agregacao(self, campo, valor):
        for r in self.filter_rows:
            f = r.get()
            if f and f["campo"] == campo and f["valor"] == valor:
                self._log(f"Filtro ja existe: {campo} = {valor}", "warn")
                return
        self._add_filter_row(campo, valor)
        self._log(f"Filtro adicionado: {campo} = {valor}. Clique ANALISAR novamente para atualizar.", "info")
        self.btn_executar.configure(state="disabled")
        self.last_total = None

    def _show_done_dialog(self, pasta, total):
        self.last_pasta = pasta
        self.progress_status.configure(text=f"Concluido. {total} tese(s) salvas.")
        self.progress_detail.configure(text=pasta)
        self._set_status(f"Concluido. Arquivos em: {pasta}")
        if messagebox.askyesno(
            "Concluido",
            f"{total} tese(s) salvas com sucesso em:\n\n{pasta}\n\nAbrir a pasta agora?",
        ):
            try:
                if sys.platform == "win32":
                    os.startfile(pasta)
                elif sys.platform == "darwin":
                    os.system(f'open "{pasta}"')
                else:
                    os.system(f'xdg-open "{pasta}"')
            except Exception as e:
                messagebox.showerror("Erro", f"Nao foi possivel abrir: {e}")

    def _log(self, text, level="info"):
        ts = datetime.now().strftime("%H:%M:%S")
        self.log_text.configure(state="normal")
        self.log_text.insert("end", f"[{ts}] ", "ts")
        self.log_text.insert("end", text + "\n", level)
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    def _clear_log(self):
        self.log_text.configure(state="normal")
        self.log_text.delete("1.0", "end")
        self.log_text.configure(state="disabled")

    def _set_status(self, text):
        self.status_label.configure(text=text)

    @staticmethod
    def _truncate(text, max_chars):
        if not text:
            return ""
        text = " ".join(text.split())  # normaliza espacos
        if len(text) <= max_chars:
            return text
        return text[: max_chars - 1] + "…"


# ============================================================
# ENTRYPOINT
# ============================================================

def main():
    app = App()
    app.mainloop()


if __name__ == "__main__":
    main()
