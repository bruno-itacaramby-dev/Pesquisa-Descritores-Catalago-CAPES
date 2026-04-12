import requests
import json
import math
import time
import os

from tqdm import tqdm
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font
from playwright.sync_api import sync_playwright


URL = "https://catalogodeteses.capes.gov.br/catalogo-teses/rest/busca"

termos = [
    "gerencialismo"
]

headers = {
    "Content-Type": "application/json",
    "User-Agent": "Mozilla/5.0"
}

keywords = []


def coletar_links(termo):

    payload = {
        "termo": termo,
        "filtros": [
            {"campo": "Grande Àrea Conhecimento", "valor": "CIÊNCIAS HUMANAS"},
            {"campo": "Área Conhecimento", "valor": "EDUCAÇÃO"},
            {"campo": "Nome Programa", "valor": "EDUCAÇÃO"},
            {"campo": "Nome Programa", "valor": "Educação"}
        ],
        "pagina": 1,
        "registrosPorPagina": 20
    }

    session = requests.Session()

    resp = session.post(URL, json=payload, headers=headers, timeout=30)
    data = resp.json()

    total = data["total"]
    por_pagina = data["registrosPorPagina"]
    total_paginas = math.ceil(total / por_pagina)

    resultados = []

    print("\nColetando páginas da API...\n")

    for pagina in tqdm(range(1, total_paginas + 1), desc="Paginas API"):

        payload["pagina"] = pagina

        resp = session.post(URL, json=payload, headers=headers, timeout=30)

        dados = resp.json()["tesesDissertacoes"]

        for item in dados:

            titulo = item.get("titulo", "")
            titulo_lower = titulo.lower()

            if not keywords or any(k in titulo_lower for k in keywords):

                data_defesa = item.get("dataDefesa")

                if data_defesa:
                    data_defesa = datetime.fromisoformat(
                        data_defesa.replace("Z", "")
                    ).strftime("%d/%m/%y")

                resultados.append({
                    "titulo": titulo,
                    "autor": item.get("autor"),
                    "grauAcademico": item.get("grauAcademico"),
                    "dataDefesa": data_defesa,
                    "link": item.get("link")
                })

        time.sleep(0.3)

    return resultados


def extrair_palavras_chave(resultados):

    print("\nAbrindo teses no Chromium...\n")

    resultados_com_resumo = []

    with sync_playwright() as p:

        browser = p.chromium.launch(headless=True)
        page = browser.new_page()

        for item in tqdm(resultados, desc="Extraindo palavras-chave"):

            try:

                page.goto(item["link"], timeout=60000)

                page.wait_for_selector("#palavras", timeout=15000)

                texto = page.locator("#palavras").inner_text().strip()

                texto = texto.rstrip(".")

                if ";" in texto:
                    partes = texto.split(";")
                else:
                    partes = texto.split(".")

                palavras = [p.strip() for p in partes if p.strip()]

                palavras_formatadas = ", ".join(palavras)

                page.wait_for_selector("#resumo", timeout=15000)

                resumo = page.locator("#resumo").inner_text().strip()

                item["palavras_chave"] = palavras_formatadas

                item_com_resumo = item.copy()
                item_com_resumo["palavras_chave"] = palavras_formatadas
                item_com_resumo["resumo"] = resumo

                resultados_com_resumo.append(item_com_resumo)

            except Exception:

                item["palavras_chave"] = ""

                item_com_resumo = item.copy()
                item_com_resumo["palavras_chave"] = ""
                item_com_resumo["resumo"] = ""

                resultados_com_resumo.append(item_com_resumo)

        browser.close()

    return resultados, resultados_com_resumo


def gerar_excel(data, termo, pasta):

    excel_file = os.path.join(pasta, f"{termo}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "dados"

    ws.append([
        "titulo",
        "autor",
        "grauAcademico",
        "dataDefesa",
        "palavras_chave",
        "resumo",
        "link"
    ])

    for item in data:

        row = ws.max_row + 1

        ws.cell(row=row, column=1, value=item["titulo"])
        ws.cell(row=row, column=2, value=item["autor"])
        ws.cell(row=row, column=3, value=item["grauAcademico"])
        ws.cell(row=row, column=4, value=item["dataDefesa"])
        ws.cell(row=row, column=5, value=item["palavras_chave"])
        ws.cell(row=row, column=6, value=item["resumo"])

        cell = ws.cell(row=row, column=7, value=item["link"])
        cell.hyperlink = item["link"]
        cell.font = Font(color="0000FF", underline="single")

    wb.save(excel_file)

    print("Excel gerado:", excel_file)


def gerar_excel_sem_resumo(data, termo, pasta):

    excel_file = os.path.join(pasta, f"{termo}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "dados"

    ws.append([
        "titulo",
        "autor",
        "grauAcademico",
        "dataDefesa",
        "palavras_chave",
        "link"
    ])

    for item in data:

        row = ws.max_row + 1

        ws.cell(row=row, column=1, value=item["titulo"])
        ws.cell(row=row, column=2, value=item["autor"])
        ws.cell(row=row, column=3, value=item["grauAcademico"])
        ws.cell(row=row, column=4, value=item["dataDefesa"])
        ws.cell(row=row, column=5, value=item["palavras_chave"])

        cell = ws.cell(row=row, column=6, value=item["link"])
        cell.hyperlink = item["link"]
        cell.font = Font(color="0000FF", underline="single")

    wb.save(excel_file)

    print("Excel gerado:", excel_file)


def salvar_json(data_sem_resumo, data_com_resumo, termo, pasta):

    json_file = os.path.join(pasta, f"{termo}.json")

    with open(json_file, "w", encoding="utf-8") as f:
        json.dump(data_sem_resumo, f, ensure_ascii=False, indent=2)

    print("JSON gerado:", json_file)

    json_file_resumo = os.path.join(pasta, f"{termo}_com_resumo.json")

    with open(json_file_resumo, "w", encoding="utf-8") as f:
        json.dump(data_com_resumo, f, ensure_ascii=False, indent=2)

    print("JSON gerado:", json_file_resumo)


def main():

    for termo in termos:

        print("\n===============================")
        print("Buscando termo:", termo)
        print("===============================\n")

        pasta = termo
        os.makedirs(pasta, exist_ok=True)

        resultados = coletar_links(termo)

        print("\nTotal de teses encontradas:", len(resultados))

        resultados_sem_resumo, resultados_com_resumo = extrair_palavras_chave(resultados)

        salvar_json(resultados_sem_resumo, resultados_com_resumo, termo, pasta)

        gerar_excel_sem_resumo(resultados_sem_resumo, termo, pasta)

        gerar_excel(resultados_com_resumo, f"{termo}_com_resumo", pasta)


if __name__ == "__main__":
    main()